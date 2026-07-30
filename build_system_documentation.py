from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from pathlib import Path
from datetime import date

OUT = Path("BroxStudies_System_Documentation.docx")
BLUE = "215B8F"
DARK = "17324D"
LIGHT = "EAF2F8"
PALE = "F4F7FA"
GRAY = "667788"
GREEN = "2D7D46"
RED = "A33A3A"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.49)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string("202830")
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for name, size, before, after, color in [
    ("Title", 30, 0, 8, DARK),
    ("Subtitle", 13, 0, 10, GRAY),
    ("Heading 1", 17, 16, 8, BLUE),
    ("Heading 2", 13.5, 12, 6, BLUE),
    ("Heading 3", 11.5, 9, 4, DARK),
]:
    st = styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = name != "Subtitle"
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for style_name in ["List Bullet", "List Number"]:
    st = styles[style_name]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.left_indent = Inches(0.38)
    st.paragraph_format.first_line_indent = Inches(-0.19)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.15

if "Callout" not in styles:
    call = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
else:
    call = styles["Callout"]
call.font.name = "Calibri"
call.font.size = Pt(10.5)
call.font.color.rgb = RGBColor.from_string(DARK)
call.paragraph_format.space_before = Pt(6)
call.paragraph_format.space_after = Pt(8)
call.paragraph_format.left_indent = Inches(0.16)
call.paragraph_format.right_indent = Inches(0.16)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(table, color="C8D2DC", size="4"):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement(f"w:{edge}")
        x.set(qn("w:val"), "single")
        x.set(qn("w:sz"), size)
        x.set(qn("w:color"), color)
        el.append(x)
    tblPr.append(el)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        shade(c, BLUE)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255,255,255)
            r.font.size = Pt(9.5)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2 == 1:
                shade(cells[i], PALE)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.size = Pt(9)
        for c in cells:
            set_cell_margins(c)
    for c in t.rows[0].cells:
        set_cell_margins(c)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
        grid = t._tbl.tblGrid
        for i, col in enumerate(grid.gridCol_lst):
            col.set(qn("w:w"), str(int(widths[i] * 1440)))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def p(text="", style=None, bold_lead=None):
    para = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r = para.add_run(bold_lead)
        r.bold = True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    return para

def bullets(items):
    for x in items:
        p(x, "List Bullet")

def numbered(items):
    for x in items:
        p(x, "List Number")

def callout(label, text, fill=LIGHT):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    t.columns[0].width = Inches(6.5)
    c = t.cell(0,0)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)
    shade(c, fill)
    set_cell_margins(c, 140, 180, 140, 180)
    borders(t, color=fill, size="0")
    para = c.paragraphs[0]
    para.style = styles["Callout"]
    r = para.add_run(label + "  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    para.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def page_break():
    doc.add_page_break()

def add_footer():
    for section in doc.sections:
        footer = section.footer
        para = footer.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        para.add_run("BroxStudies System Documentation  |  ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        para._p.append(fld)
        for r in para.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string(GRAY)

# Cover
p("BROXSTUDIES", "Subtitle")
p("System Documentation", "Title")
p("Features, user workflows, architecture, APIs, operations, security, and maintenance", "Subtitle")
doc.add_paragraph("\n")
table(
    ["Document field", "Value"],
    [
        ("System", "BroxStudies Online - AI-powered study and examination platform"),
        ("Codebase version", "Application version 0.1.0 (repository snapshot)"),
        ("Documentation date", "29 July 2026"),
        ("Audience", "Product owners, administrators, developers, support staff, and technical partners"),
        ("Source of truth", "Implemented frontend and backend code in this repository"),
    ],
    [1.75, 4.75],
)
callout("Purpose", "This guide explains what the system does, who uses each capability, how the parts work together, how to run and configure it, and where operational or security attention is required.")
p("Document scope", "Heading 2")
bullets([
    "Covers the active React frontend, FastAPI backend, persistence, integrations, background services, deployment files, and principal user journeys.",
    "Distinguishes implemented features from optional integrations and roadmap recommendations.",
    "Does not reproduce secret values, live credentials, private user data, or payment identifiers.",
])
page_break()

# TOC
p("Contents", "Heading 1")
toc = [
("1", "Executive overview"), ("2", "Users, access, and subscriptions"),
("3", "Feature catalogue"), ("4", "Student experience and workflows"),
("5", "Administrator and teacher capabilities"), ("6", "System architecture"),
("7", "AI, RAG, content, and question generation"), ("8", "Data and persistence"),
("9", "API reference"), ("10", "Configuration and integrations"),
("11", "Local development and deployment"), ("12", "Operations, testing, and troubleshooting"),
("13", "Security and privacy"), ("14", "Known constraints and recommended improvements"),
("A", "Frontend route map"), ("B", "Backend endpoint inventory"), ("C", "Glossary"),
]
table(["Section", "Topic"], toc, [0.7, 5.8])
callout("Navigation note", "Word can generate a dynamic table of contents from the Heading styles. The table above is a stable reading map and remains useful in viewers that do not update Word fields.")

page_break()
p("1. Executive overview", "Heading 1")
p("BroxStudies is a responsive web application for Ghanaian secondary-school study and exam preparation. It combines curriculum resources, past-question intelligence, AI-assisted tutoring and question generation, practice and mock-exam workflows, progress analytics, learning plans, competitions, and subscription-controlled access in one platform.")
p("At runtime, a React single-page application communicates with a FastAPI service through JSON and multipart HTTP endpoints. The backend manages identity, progress, access codes, payments, document processing, book content, AI calls, question generation, analytics, and administrative content. In production, FastAPI also serves the compiled frontend.")
p("System outcomes", "Heading 2")
table(["Outcome", "How the system supports it"], [
("Study", "Topic-based study, curriculum resources, book library, news, and guided learning plans."),
("Practise", "Generated questions, question banks, timed quizzes, marking, answer explanations, and history."),
("Prepare for WASSCE", "Past-question extraction, pattern analysis, likely-question generation, structured papers, and marking."),
("Get help", "Streaming AI tutor chat, image interpretation, file-assisted questions, mathematics rendering, and chat history."),
("Measure progress", "Mastery records, analytics, streaks, XP/gamification, mock results, rankings, and review scheduling."),
("Operate the service", "Admin analytics, coupons/access codes, payment review, SMS logs, competitions, news, and social content."),
], [1.55,4.95])
p("Key architectural characteristics", "Heading 2")
bullets([
"Single deployable web service: the backend API can serve the built SPA and uploaded assets.",
"Provider-flexible LLM configuration: OpenAI-compatible, DeepSeek, and generic LLM settings with fallback behavior.",
"Retrieval-augmented generation: uploaded and curated documents are chunked and indexed for context-aware generation.",
"Local-first development persistence with SQLite, plus configurable PostgreSQL for hosted environments.",
"Progressive loading and background refresh for large content collections and external news.",
"Installable PWA capabilities, offline history/support, push subscriptions, lazy-loaded routes, and responsive UI components.",
])
callout("Implementation status", "The document describes code present in the repository. Some features depend on environment variables, external provider accounts, uploaded content, or administrator configuration before they become operational.")

p("2. Users, access, and subscriptions", "Heading 1")
table(["Role/state", "Capabilities and restrictions"], [
("Guest", "Can reach public/landing flows and limited study entry points; gated screens prompt sign-in or subscription."),
("Registered student", "Owns progress, tutor history, learning profile, review cards, exams, classes, and settings."),
("Subscribed/activated student", "Receives access to subscription-gated practice, WASSCE, quiz, news, rankings, library, history, and analytics experiences."),
("Teacher", "Can create classes, issue assignments, review teacher snapshots, and use learning-management endpoints."),
("Administrator", "Uses a protected admin route for analytics, codes/coupons, payments, SMS, competitions, news, and social content."),
], [1.55,4.95])
p("Authentication lifecycle", "Heading 2")
numbered([
"A user signs up with account details or signs in using password-based authentication or Google identity.",
"Where required, the system sends and verifies a one-time password (OTP). Phone numbers can be added and separately verified.",
"A bearer access token represents the signed-in session; the frontend retrieves the current user and subscription state.",
"Forgot-password endpoints request and confirm a reset. Users can update progress and delete their account.",
"Subscription can be activated with an access code or through a supported payment flow, depending on deployment configuration.",
])
p("Subscription and activation", "Heading 2")
bullets([
"Access-code verification activates entitlement associated with a generated code.",
"Administrators can generate code batches/coupons, inspect inventory, and send codes by SMS.",
"Moolre supports mobile-money initiation, OTP submission, status checks, transaction history, and webhooks.",
"Paystack initialization, verification, and webhook routes remain available as a configurable fallback.",
"A manual payment-request and administrator confirmation/rejection path supports assisted operations.",
])
callout("Security warning", "The repository contains default development values for administrative and signing settings. Production deployments must supply strong secrets through environment variables, rotate any exposed defaults, restrict administrator access, and verify webhook signatures.")

p("3. Feature catalogue", "Heading 1")
p("3.1 Study dashboard and navigation", "Heading 2")
bullets([
"Dashboard summary with quick actions, progress signals, alerts, promotional information, and navigation to core study areas.",
"Academic-track selection used to tailor levels, subjects, and curriculum choices.",
"Responsive application shell with desktop/mobile navigation, loading gates, page transitions, theme support, and account controls.",
"PWA installation support, push handler, offline-aware history utilities, and local storage for selected client-side progress.",
])
p("3.2 AI tutor", "Heading 2")
bullets([
"Ask text questions and receive a structured answer with sources/context where available.",
"Stream answers for a more immediate chat experience.",
"Upload supporting files with a tutor request.",
"Interpret an image, useful for photographed questions, diagrams, or handwritten work.",
"Persist and retrieve user chat history; guests can retain limited local chat history.",
"Render mathematical notation through KaTeX and Markdown/GFM content.",
"Track AI usage so deployments can apply quotas and monitor consumption.",
])
p("3.3 Question generation and practice", "Heading 2")
bullets([
"Generate questions by academic level, subject, type, count, difficulty, topics, and source selection.",
"Support multiple-choice, short-answer, essay, structured, and true/false-style formats where the generation path permits.",
"Run generation as a job and poll job status; view current resource and loading status.",
"Mark practice submissions and return per-item scoring, totals, explanations, and feedback.",
"Generate professional/standardized outputs and downloadable question PDFs.",
"Grade uploaded answer PDFs using the dedicated grading flow.",
"Save and retrieve exam-history records.",
])
p("3.4 WASSCE intelligence", "Heading 2")
bullets([
"Extract and organize historical past questions by subject, paper, section, year, topic, and question structure.",
"Analyze frequency, trends, topic distribution, gaps, recency, confidence, and likely examination patterns.",
"Generate likely WASSCE questions using specialized services, curriculum sources, past-question archives, and configured LLMs.",
"Build paper-like outputs that distinguish objective, theory, structured, and subject-specific section rules.",
"Attach solutions/explanations when available and export/share generated material.",
])
p("3.5 Live quiz", "Heading 2")
bullets([
"Create a quiz room and receive a join code.",
"Participants join using the code and retrieve current room state.",
"The host starts the session; participants submit answers.",
"The service calculates and returns quiz state/results for the shared session.",
])
p("3.6 Learning hub", "Heading 2")
bullets([
"Maintain an academic learning profile and mastery observations by subject and topic.",
"Run diagnostics and generate an adaptive revision plan; mark plan items complete.",
"Record mock-exam performance and expose an overview of the student’s plan and mastery.",
"Create spaced-review cards, list cards due for review, and grade recall quality to schedule future reviews.",
"Download an offline learning pack.",
"Join classes, receive assignments, and support teacher-owned class creation.",
"Submit feedback, appeals, post reports, and user blocks; moderation endpoints manage reports.",
"Subscribe to push notifications and send configured notifications.",
])
p("3.7 Library and source content", "Heading 2")
bullets([
"Search a curated book/resource library and open individual records.",
"Retrieve excerpts and generate a quiz based on selected book content.",
"Upload PDFs as syllabi, past questions, or textbooks and inspect processing status.",
"Load deferred documents and selectively rebuild or refresh indexed resources.",
"Fetch curriculum resources, monitor/cancel fetching, and list available resources.",
"Use Source Studio to work with selected source material and generation inputs.",
])
p("3.8 Analytics, history, and gamification", "Heading 2")
bullets([
"Student analytics views combine practice activity, accuracy, subject performance, weekly activity, and progress signals.",
"History collects completed/generated examinations and supports re-opening or exporting past work.",
"Gamification utilities manage XP, levels, streaks, achievements, celebrations, and progress feedback.",
"Rankings and competition leaderboards provide comparative engagement views.",
])
p("3.9 News, competitions, and community content", "Heading 2")
bullets([
"News view combines administrator-managed articles with optional background-fetched external stories.",
"Competitions can be created, supplied with PDFs/images, opened for registration, and displayed with leaderboards.",
"Administrator-managed social posts support comments and reactions.",
"Startup Cup promotional and voting components provide campaign-specific calls to action.",
])
p("3.10 Settings and exports", "Heading 2")
bullets([
"Account settings, theme preferences, profile/phone management, and account deletion.",
"Client-side export utilities produce printable/downloadable study-question material.",
"Professional PDF generation is also available through the backend for consistent server-produced outputs.",
])

page_break()
p("4. Student experience and workflows", "Heading 1")
p("First-time onboarding", "Heading 2")
numbered([
"Open the landing page and create an account or sign in.",
"Complete OTP or Google authentication when configured.",
"Select the academic track/level that determines available subjects.",
"Add and verify a phone number if SMS or payment activation requires it.",
"Activate a subscription using an access code or complete the enabled payment flow.",
"Open the dashboard and choose Study, Tutor, Practice, WASSCE, Learning, Library, or another module.",
])
p("Generate and complete a practice set", "Heading 2")
numbered([
"Choose academic level, subject, question type, topics, difficulty, and number of questions.",
"The frontend submits a generation request; long-running work may return a job identifier.",
"The backend selects curriculum/past-question context, calls the configured generator, and normalizes question structures.",
"The student answers questions in the practice or exam interface.",
"The marking endpoint calculates results and returns feedback/explanations.",
"The result can update local gamification, server progress, mastery, analytics, and exam history.",
])
p("Ask the AI tutor", "Heading 2")
numbered([
"Open Study/Tutor and enter a question, optionally adding a file or image.",
"The system retrieves relevant source context when available and sends a guarded prompt to the selected model provider.",
"The response streams or returns as a completed answer and is rendered with Markdown and mathematics support.",
"Authenticated chat history is retained; usage counters help enforce platform limits.",
])
p("Build an adaptive plan", "Heading 2")
numbered([
"Save the learner profile and record or import mastery observations.",
"Complete a diagnostic or mock examination.",
"Generate a plan from weaker topics and current learning state.",
"Complete plan items and create review cards from study material.",
"Review due cards and grade recall; the service updates future review scheduling.",
])
p("Prepare with WASSCE material", "Heading 2")
numbered([
"Choose the academic track and subject.",
"Review topic and pattern information derived from historical sources.",
"Request likely questions or a standardized paper.",
"Complete, mark, save, or export the result, then use analytics to identify weak areas.",
])

p("5. Administrator and teacher capabilities", "Heading 1")
p("Administrator workspace", "Heading 2")
table(["Area", "Principal actions"], [
("Analytics", "Review user, subscription, usage, payment, and engagement summaries exposed by the admin analytics model."),
("Access codes", "Inspect inventory, generate coupons/codes, and send selected codes by SMS."),
("Payments", "Review pending manual requests, confirm/reject them, and inspect provider transaction history."),
("Competitions", "Create competitions, upload PDF/image assets, register entries, and view leaderboards."),
("News", "Create, update, delete, image-enable, and list public or draft news articles."),
("Social", "Create posts, accept comments, and update reactions."),
("Messaging", "Review SMS delivery logs and provider-related outcomes."),
("Moderation", "Review reported community content and resolve moderation records through learning endpoints."),
], [1.35,5.15])
p("Teacher workflows", "Heading 2")
bullets([
"Create a class and share its join code.",
"List teacher-owned classes and assignments.",
"Create assignments for a class.",
"View an aggregate teacher snapshot for supported learner activity.",
"Use generated questions, library content, and professional PDFs as instructional materials.",
])
callout("Authorization boundary", "The UI route gates improve user experience, but backend authorization is the real control boundary. Administrator, teacher, subscription, and user ownership checks must be enforced on every corresponding API operation.")

p("6. System architecture", "Heading 1")
p("Logical architecture", "Heading 2")
table(["Layer", "Responsibilities", "Primary technology"], [
("Browser/PWA", "Navigation, forms, study UI, local/offline state, visualizations, exports, push handling.", "React 18, TypeScript, Vite, Tailwind, Radix UI"),
("API", "Routing, validation, authentication dependencies, responses, static SPA delivery.", "FastAPI, Pydantic"),
("Domain services", "Auth, learning, question generation, WASSCE intelligence, tutoring, payments, SMS, PDF processing.", "Python services"),
("AI/RAG", "Prompt orchestration, retrieval, embeddings, provider selection, structured generation.", "OpenAI-compatible APIs, DeepSeek, LangChain, FAISS"),
("Persistence", "Users, tokens, progress, access, learning records, content metadata, transactions.", "SQLite locally; PostgreSQL configurable"),
("Files/content", "Uploaded PDFs, competition assets, curated curriculum, vector stores, generated PDFs.", "Filesystem + application directories"),
("External services", "LLM provider, Google identity, Moolre SMS/payments, optional Paystack, web push, external news.", "HTTPS APIs/webhooks"),
], [1.15,3.2,2.15])
p("Request lifecycle", "Heading 2")
numbered([
"React calls a typed function in frontend/src/api/endpoints.ts through the shared Axios client.",
"FastAPI matches a versionless /api route and validates the request using Pydantic models.",
"Authentication/authorization resolves the current user when the route is protected.",
"A route delegates substantive work to a domain service or background task.",
"The service reads persistence/content, optionally performs retrieval or invokes an external provider, and normalizes results.",
"FastAPI returns JSON, a file, or a stream; React updates the screen and relevant local/server progress.",
])
p("Startup and background work", "Heading 2")
bullets([
"Ensure content and persistence directories exist.",
"Generate the subject catalogue if it is missing.",
"Optionally preload past-question caches.",
"Start external-news refresh when enabled.",
"Auto-load configured document collections in the background and expose loading status.",
"Serve static frontend assets and fall back to index.html for client-side routes.",
])
p("Repository structure", "Heading 2")
table(["Path", "Purpose"], [
("backend/app/main.py", "FastAPI application, lifecycle, middleware, router registration, static delivery."),
("backend/app/routes/", "HTTP endpoints grouped by domain."),
("backend/app/services/", "Business logic, provider clients, retrieval, content processing, generation."),
("backend/app/models.py", "Request and response contracts."),
("backend/data/ and vector_store/", "Local persistence, curated resources, caches, and vector indexes."),
("frontend/src/pages/", "Top-level routed screens."),
("frontend/src/components/", "Shared shell, auth, chat, exam, charts, dashboard, and UI components."),
("frontend/src/hooks/", "Authentication, gamification, analytics, alerts, PWA, theme, generation, offline state."),
("frontend/src/api/", "HTTP client, endpoint wrappers, and TypeScript data types."),
("Dockerfile / render.yaml", "Container build and Render deployment configuration."),
("start_system.ps1 / .bat", "Windows local-development launch helpers."),
], [2.15,4.35])

p("7. AI, RAG, content, and question generation", "Heading 1")
p("Content pipeline", "Heading 2")
numbered([
"Documents are uploaded or discovered in configured syllabus, past-question, textbook, or resource locations.",
"The PDF processor extracts text. Batch loaders categorize and process collections without blocking application startup.",
"Text is divided using configured chunk size and overlap.",
"Embeddings are created with the configured embedding model and stored/retrieved through the vector-store layer.",
"At generation time, subject/topic queries retrieve relevant chunks and combine them with past-question intelligence.",
"The generation service prompts the selected LLM, parses/normalizes structured outputs, and returns validated question objects.",
])
p("Generation modes", "Heading 2")
table(["Mode", "Purpose"], [
("Standard generation", "General questions by subject, topic, type, count, and difficulty."),
("Professional generation", "More controlled/standardized question sets suitable for formatted delivery."),
("Likely WASSCE generation", "Exam-oriented output informed by archive patterns, syllabus constraints, and section rules."),
("Tutor response", "Explanatory answer grounded in retrieved or attached context."),
("Book quiz", "Questions generated from a selected book or excerpt."),
("Live quiz", "A shared session built from generated or selected questions."),
("PDF grading", "Interpret and grade answers submitted in PDF form."),
], [1.75,4.75])
p("Provider selection and resilience", "Heading 2")
bullets([
"Explicit generic LLM settings take priority when provided.",
"Otherwise, OpenAI configuration is used when an OpenAI key is present.",
"DeepSeek can be selected when its key is available and OpenAI is not configured.",
"Fallback behavior is controlled by an environment flag.",
"Generation should fail with actionable provider/configuration errors rather than expose secrets or raw prompts.",
])
callout("Quality principle", "Generated material is probabilistic. Examination content, solutions, marking, and predictions should be reviewed by a qualified educator before high-stakes use. Pattern confidence is evidence, not a guarantee of future examination content.")

p("8. Data and persistence", "Heading 1")
table(["Data class", "Examples", "Storage/handling"], [
("Identity", "Users, password hashes, phones, OTP state, Google identifiers, tokens.", "Authentication database via SQLite/PostgreSQL."),
("Entitlement", "Subscriptions, access codes, coupons, payment requests and provider transactions.", "Authentication/business persistence."),
("Learning", "Profiles, mastery, plans, mocks, review cards, classes, assignments.", "Learning tables created/managed by the learning service."),
("Activity", "Exam history, progress, tutor chat history, AI usage.", "User-linked persistence and selected client-side caches."),
("Content", "Syllabi, past questions, textbooks, books, subjects catalogue.", "Filesystem/catalogue plus indexed/vector representations."),
("Editorial", "News, competitions, leaderboard entries, social posts/comments/reactions.", "Application persistence and uploaded assets."),
("Operational", "SMS logs, loading manifests, caches, provider references.", "Database/files/logs depending on subsystem."),
], [1.15,2.35,3.0])
p("Database portability", "Heading 2")
p("When DATABASE_URL is absent, the application resolves to a local SQLite database under the persistence directory. PostgreSQL URLs are normalized for common hosting formats, including the postgres:// alias. Hosted deployments should use managed PostgreSQL, backups, migrations, connection security, and environment-scoped credentials.")
p("File lifecycle", "Heading 2")
bullets([
"PDF uploads are stored under the configured upload directory.",
"Generated vector data is stored under the configured vector-store directory.",
"Curated subject resources live in the site-resource/data locations.",
"Competition uploads are exposed under /uploads; production deployments should validate type/size and consider object storage.",
"Generated PDFs are returned as downloadable artifacts or temporary files according to the service path.",
])
p("Backup and retention recommendations", "Heading 2")
bullets([
"Back up the production database and uploaded/curated content independently.",
"Define retention for OTPs, tokens, tutor messages, exam history, uploaded answer files, SMS logs, and payment events.",
"Do not commit live .env files, database snapshots, uploaded personal documents, or provider responses.",
"Use migrations and restore tests before schema changes or deployment moves.",
])

p("9. API reference", "Heading 1")
p("All application endpoints use the /api prefix except the root SPA, health check, uploaded assets, and the frontend fallback. Interactive OpenAPI documentation is configured at /api/docs and ReDoc at /api/redoc.")
table(["API group", "Base path", "Responsibility"], [
("Authentication", "/api/auth", "Identity, OTP, profile, progress, subscription, activation, payment request."),
("Tutor", "/api/tutor", "AI questions, streaming, images/files, history, usage."),
("Questions", "/api/questions", "Generation, jobs, marking, PDF grading, quizzes, history, PDF output."),
("Analysis", "/api/analysis", "Patterns, topics, and vector-index rebuild."),
("Uploads", "/api/uploads", "PDF intake, status, and deferred loading."),
("Resources", "/api/resources", "Curriculum fetching, status, cancellation, availability."),
("Books", "/api/books", "Search, details, excerpt, book-based quiz."),
("Learning", "/api/learning", "Mastery, plans, reviews, classes, push, feedback, moderation."),
("Payments", "/api/payments", "Moolre and Paystack initiation, verification, history, webhooks."),
("Admin", "/api/admin", "Operational administration and editorial management."),
], [1.35,1.65,3.5])
callout("Contract note", "The complete endpoint inventory appears in Appendix B. Request/response field definitions remain authoritative in backend/app/models.py and the generated OpenAPI schema.")

p("10. Configuration and integrations", "Heading 1")
table(["Configuration area", "Important variables (names only)"], [
("LLM", "OPENAI_API_KEY, OPENAI_BASE_URL, OPENAI_MODEL, DEEPSEEK_API_KEY, DEEPSEEK_BASE_URL, DEEPSEEK_MODEL, LLM_*"),
("Persistence", "DATABASE_URL, PERSISTENCE_DIR, AUTH_DB_PATH"),
("Authentication", "AUTH_SECRET_KEY, ACCESS_TOKEN_EXPIRE_HOURS, GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET, OTP_EXPIRE_MINUTES"),
("Administration", "ADMIN_USERNAME, ADMIN_PASSWORD, ADMIN_SECRET"),
("Subscription", "SUBSCRIPTION_PRICE_GHS, SUBSCRIPTION_MONTHS"),
("Moolre", "MOOLRE_VAS_KEY, MOOLRE_SENDER_ID, MOOLRE_API_*, MOOLRE_ACCOUNT_NUMBER, MOOLRE_PAYMENT_ENABLED, SMS_ENABLED"),
("Paystack", "PAYSTACK_SECRET_KEY, PAYSTACK_PUBLIC_KEY, PAYSTACK_ENABLED"),
("Web/PWA", "PUBLIC_APP_URL, CORS_ORIGINS, VAPID_PUBLIC_KEY, VAPID_PRIVATE_KEY, VAPID_CONTACT_EMAIL"),
("Content", "PDF_UPLOAD_DIR, DATA_DIR, SITE_RESOURCE_DIR, CHUNK_SIZE, CHUNK_OVERLAP, VECTOR_STORE_DIR, EMBEDDING_MODEL"),
("Generation", "MAX_QUESTIONS, MIN_CONFIDENCE_SCORE"),
("Startup", "AUTO_LOAD_ON_STARTUP, PRELOAD_PAST_QUESTIONS_ON_STARTUP, NEWS_REFRESH_ENABLED, LAZY_LOAD, SELECTIVE_LOAD"),
], [1.35,5.15])
p("External integration behavior", "Heading 2")
table(["Integration", "What it enables", "Operational requirement"], [
("LLM provider", "Tutor and question-generation intelligence.", "Valid API credentials, model access, quotas, timeout/cost monitoring."),
("Google Identity", "Google sign-in.", "OAuth client, allowed origins/redirect configuration."),
("Moolre SMS", "OTP and access-code delivery.", "VAS key, approved sender, phone normalization, delivery monitoring."),
("Moolre Payments", "Mobile-money collection and OTP confirmation.", "API account, secure webhooks, reconciliation."),
("Paystack", "Alternative online payment path.", "Secret/public keys and verified webhook configuration."),
("Web Push", "Browser notifications.", "VAPID key pair, HTTPS, user consent."),
("External News", "Background news refresh.", "Network availability and source/error handling."),
], [1.35,2.35,2.8])

p("11. Local development and deployment", "Heading 1")
p("Prerequisites", "Heading 2")
bullets(["Python 3.9 or newer.", "Node.js 18 or newer with npm.", "At least one configured LLM provider for AI functionality.", "Optional provider credentials for Google, SMS, payments, and push."])
p("Recommended local start", "Heading 2")
numbered([
"Copy backend/.env.example to backend/.env and supply development-only values.",
"Run .\\start_system.ps1 from the repository root. Use -SkipInstall, -BackendOnly, or -FrontendOnly when appropriate.",
"Alternatively install backend requirements and start uvicorn app.main:app --reload from backend.",
"Install frontend dependencies and run npm run dev from frontend.",
"Open the Vite development URL; the API normally runs at http://localhost:8000.",
])
p("Production build", "Heading 2")
numbered([
"Install frontend dependencies and run the TypeScript/Vite production build.",
"Install Python requirements in the runtime image.",
"Provide production environment variables through the hosting secret manager.",
"Start the FastAPI application so it serves API routes and frontend/dist.",
"Provision durable database and file/object storage as required; do not rely on ephemeral disk for irreplaceable data.",
])
p("Deployment assets", "Heading 2")
bullets([
"Dockerfile defines the unified container build.",
"render.yaml defines the Render.com service configuration.",
"The backend mounts compiled assets and uploaded files and supplies SPA history fallback.",
"Health monitoring can call GET /health; document-loading status is available at GET /api/status/documents.",
])

p("12. Operations, testing, and troubleshooting", "Heading 1")
p("Routine operational checks", "Heading 2")
table(["Check", "Expected signal"], [
("Health", "GET /health returns a healthy status."),
("Frontend", "Root URL loads the SPA and client-side routes survive refresh."),
("Document loading", "/api/status/documents and question loading-progress endpoints move to complete."),
("Authentication", "Signup/login/OTP produce a valid session without exposing verification data."),
("AI", "A small tutor and question-generation request succeeds within quota and timeout."),
("Payments", "Provider sandbox transaction and signed webhook reconcile to one subscription event."),
("SMS", "Delivery logs show accepted/sent outcomes without logging OTP or secret data."),
("Persistence", "Restart preserves accounts, learning state, and admin content."),
], [1.55,4.95])
p("Automated checks", "Heading 2")
bullets([
"Backend: run pytest backend/ (the repository includes authentication normalization, question generation, SMS, structured-question, and WASSCE-related tests).",
"Frontend: run npm run typecheck --prefix frontend, npm run test --prefix frontend, and npm run build --prefix frontend.",
"Provider parity: use the included LLM parity test only with safe test credentials and controlled spend.",
"Before release, test guest, student, subscribed student, teacher, and administrator permission boundaries.",
])
p("Troubleshooting", "Heading 2")
table(["Symptom", "Likely checks"], [
("AI says provider is unavailable", "Confirm selected provider key/model/base URL; inspect fallback setting and application logs."),
("PDF upload fails", "Validate PDF, size/type limits, upload-directory permissions, disk capacity, and extraction errors."),
("Questions remain loading", "Inspect background loader status, manifests, data paths, memory, and lazy/selective-load flags."),
("Frontend cannot reach API", "Check API base URL/proxy, server port, CORS origins, HTTPS mixed-content rules, and browser console."),
("Login/OTP fails", "Check signing secret, token clock/expiry, SMS configuration, phone normalization, and database writes."),
("Payment remains pending", "Check provider status, webhook delivery/signature, external reference, idempotency, and admin reconciliation."),
("Refresh returns 404", "Ensure compiled frontend exists and FastAPI SPA fallback/static configuration is active."),
("Data disappears after deploy", "Verify persistent database/object storage; hosted ephemeral filesystems are not durable."),
], [1.8,4.7])

p("13. Security and privacy", "Heading 1")
p("Required production controls", "Heading 2")
bullets([
"Replace every default secret and administrative credential; keep secrets only in the deployment secret manager.",
"Hash passwords with a modern adaptive algorithm and never log passwords, OTPs, bearer tokens, provider keys, or full payment payloads.",
"Enforce role, subscription, and resource-ownership checks in backend routes, not only frontend gates.",
"Validate upload MIME type, file signature, filename, size, page count, and extraction limits; scan files if feasible.",
"Rate-limit login, OTP, password reset, tutor, generation, upload, and payment endpoints.",
"Verify Moolre/Paystack webhook authenticity, use replay protection and idempotent transaction processing.",
"Use HTTPS, restrictive production CORS, secure headers, short-lived tokens where practical, and safe token storage.",
"Minimize collection of phone numbers, chat content, answer files, and learning data; publish retention/deletion rules.",
"Protect generated/admin file paths from traversal and unauthorized enumeration.",
"Audit administrator actions and sensitive entitlement changes.",
])
p("Privacy-sensitive data", "Heading 2")
table(["Data", "Primary risk", "Recommended treatment"], [
("Student identity/phone", "Unauthorized disclosure or account takeover.", "Encrypt in transit, restrict access, normalize safely, minimize display."),
("Tutor chats/files/images", "Sensitive academic or personal content.", "Consent, retention controls, access isolation, safe provider disclosure."),
("Learning analytics", "Profiling of minors and performance.", "Purpose limitation, guardian/school policy where applicable, export/deletion."),
("Payment records", "Financial fraud and privacy exposure.", "Store provider references and minimum metadata; avoid raw sensitive payment data."),
("OTP/tokens", "Session compromise.", "Hash or protect, expire quickly, one-time use, redact from logs."),
], [1.45,2.2,2.85])
callout("Immediate review item", "Treat any default administrator credentials, signing secrets, test databases, .env files, or real provider values found in source control as potentially exposed. Rotate/remove them and review repository history before production use.", "FCE8E8")

p("14. Known constraints and recommended improvements", "Heading 1")
table(["Priority", "Recommendation", "Why it matters"], [
("Critical", "Remove/rotate defaults and committed secrets; add secret scanning.", "Prevents account, payment, messaging, and token compromise."),
("Critical", "Formalize RBAC and endpoint-level authorization tests.", "Frontend gates alone cannot protect data or admin actions."),
("High", "Add database migrations, managed PostgreSQL, backups, and restore drills.", "Reduces schema drift and data-loss risk."),
("High", "Harden uploads and move durable files to object storage.", "Improves security, scaling, and hosted durability."),
("High", "Add rate limits, job queue, timeouts, retries, and idempotency.", "Protects expensive AI/payment paths and improves reliability."),
("High", "Document privacy, consent, retention, and deletion for student data.", "The platform may process minors’ identity and performance data."),
("Medium", "Expand API integration, browser journey, and permission tests.", "Current tests cover selected services but not the full product surface."),
("Medium", "Introduce structured logs, metrics, tracing, and alerting.", "Makes provider failures and slow generation diagnosable."),
("Medium", "Version the API and publish stable request/response examples.", "Supports future mobile/LMS clients and safer evolution."),
("Medium", "Add educator review and provenance indicators to AI outputs.", "Reduces overreliance on generated predictions or marking."),
("Low", "Reconcile README with current implementation each release.", "The current product surface is broader than the summary documentation."),
], [0.75,3.45,2.3])
p("Roadmap opportunities", "Heading 2")
bullets([
"Native mobile clients or deeper offline synchronization.",
"LMS interoperability, roster sync, and assignment standards.",
"Instructor authoring/review workflows and collaborative question banks.",
"Additional document formats and safer document-conversion pipeline.",
"More transparent model citations, evaluation sets, and subject-specific quality dashboards.",
])

page_break()
p("Appendix A. Frontend route map", "Heading 1")
frontend_routes = [
("/", "Study/Tutor home"), ("/welcome", "Public landing"), ("/login", "Sign in"), ("/signup", "Register"),
("/forgot-password", "Password recovery"), ("/activate", "Subscription/access activation"), ("/select-track", "Academic track selection"),
("/dashboard", "Student dashboard"), ("/docs", "Documentation/help content"), ("/learning", "Adaptive learning hub"),
("/source-studio", "Source-driven generation workspace"), ("/admin", "Administrator workspace"), ("/analytics", "Student analytics"),
("/practice", "Practice questions"), ("/wassce", "WASSCE intelligence and papers"), ("/quiz", "Live/shared quiz"),
("/news", "Education/news content"), ("/rankings", "Rankings/leaderboards"), ("/library", "Book/resource library"),
("/history", "Exam/generation history"), ("/settings", "Account and application settings"), ("/vote", "Startup Cup redirect"),
]
table(["Route", "Screen/purpose"], frontend_routes, [1.75,4.75])
p("Most application-shell routes are additionally wrapped in sign-in and/or subscription gates. The exact gate configuration is defined in frontend/src/App.tsx.")

p("Appendix B. Backend endpoint inventory", "Heading 1")
endpoint_groups = {
"System": ["GET /health", "GET /api/status/documents", "GET /api/docs", "GET /api/redoc"],
"Uploads": ["POST /api/uploads/pdf", "GET /api/uploads/status", "POST /api/uploads/load-deferred"],
"Tutor": ["GET /api/tutor/usage", "POST /api/tutor/ask", "POST /api/tutor/ask/stream", "POST /api/tutor/interpret-image", "POST /api/tutor/ask-with-files", "GET /api/tutor/history"],
"Resources": ["POST /api/resources/fetch-curriculum-resources", "GET /api/resources/fetch-curriculum-resources/status", "GET /api/resources/available-resources", "POST /api/resources/fetch-curriculum-resources/cancel"],
"Questions": ["GET /api/questions/loading-progress", "POST /api/questions/load-remaining", "POST /api/questions/generate", "GET /api/questions/jobs/{job_id}", "GET /api/questions/jobs", "GET /api/questions/subjects", "GET /api/questions/question-types", "GET /api/questions/resource-status", "POST /api/questions/mark-practice", "POST /api/questions/grade-answers-pdf", "POST /api/questions/quiz/create", "POST /api/questions/quiz/join", "GET /api/questions/quiz/{code}/state", "POST /api/questions/quiz/{code}/start", "POST /api/questions/quiz/{code}/submit", "POST /api/questions/history/exams", "GET /api/questions/history/exams", "POST /api/questions/generate-pdf", "POST /api/questions/generate-professional"],
"Analysis": ["GET /api/analysis/patterns/{subject}", "GET /api/analysis/topics/{subject}", "POST /api/analysis/index"],
"Books": ["GET /api/books/search", "GET /api/books/{book_id}", "POST /api/books/quiz", "GET /api/books/{book_id}/excerpt"],
"Authentication": ["GET /api/auth/config", "POST /api/auth/signup", "POST /api/auth/login", "POST /api/auth/google", "POST /api/auth/request-otp", "POST /api/auth/verify-otp", "POST /api/auth/add-phone", "POST /api/auth/verify-phone", "GET /api/auth/me", "POST /api/auth/forgot-password/request", "POST /api/auth/forgot-password/confirm", "DELETE /api/auth/account", "GET /api/auth/progress", "PATCH /api/auth/progress", "GET /api/auth/subscription", "POST /api/auth/verify-code", "POST /api/auth/payment-request", "POST /api/auth/admin/generate-codes"],
"Payments": ["POST /api/payments/moolre/initiate", "POST /api/payments/moolre/submit-otp", "GET /api/payments/moolre/status/{external_ref}", "GET /api/payments/moolre/history", "POST /api/payments/moolre/webhook", "POST /api/payments/paystack/initialize", "GET /api/payments/paystack/verify/{reference}", "POST /api/payments/paystack/webhook"],
"Learning": ["GET /api/learning/overview", "PUT /api/learning/profile", "POST /api/learning/mastery", "POST /api/learning/diagnostic", "POST /api/learning/plan/generate", "PUT /api/learning/plan/{item_id}/complete", "POST /api/learning/mock/complete", "GET /api/learning/offline-pack", "POST /api/learning/reviews/cards", "GET /api/learning/reviews/due", "PUT /api/learning/reviews/{card_id}/grade", "POST /api/learning/classes/join", "GET /api/learning/classes", "POST /api/learning/moderation/appeals", "GET /api/learning/push/config", "POST /api/learning/push/subscribe", "POST /api/learning/push/send", "POST /api/learning/feedback", "POST /api/learning/community/posts/{post_id}/report", "POST /api/learning/community/users/{blocked_id}/block", "GET /api/learning/teacher/snapshot", "POST /api/learning/teacher/classes", "GET /api/learning/teacher/classes", "POST /api/learning/teacher/classes/{class_id}/assignments", "GET /api/learning/moderation/reports", "PUT /api/learning/moderation/reports/{report_id}"],
"Admin": ["POST /api/admin/login-secret", "GET /api/admin/analytics", "POST /api/admin/competitions", "POST /api/admin/competitions/{comp_id}/upload-pdf", "POST /api/admin/competitions/{comp_id}/upload-image", "GET /api/admin/coupons/inventory", "POST /api/admin/coupons/generate", "GET /api/admin/payments/pending", "POST /api/admin/payments/{request_id}/confirm", "POST /api/admin/payments/{request_id}/reject", "POST /api/admin/codes/send-sms", "GET /api/admin/sms-log", "GET /api/admin/competitions/all", "GET /api/admin/competitions", "POST /api/admin/competitions/{comp_id}/register", "GET /api/admin/leaderboard", "GET /api/admin/news", "GET /api/admin/news/all", "POST /api/admin/news", "PUT /api/admin/news/{article_id}", "DELETE /api/admin/news/{article_id}", "POST /api/admin/news/{article_id}/upload-image", "GET /api/admin/social", "POST /api/admin/social", "POST /api/admin/social/{post_id}/comments", "PUT /api/admin/social/{post_id}/reaction"],
}
for group, endpoints in endpoint_groups.items():
    p(group, "Heading 2")
    table(["Method and path", "Purpose"], [(e, {
        "GET":"Retrieve or inspect", "POST":"Create, submit, or trigger", "PUT":"Replace/update", "PATCH":"Partially update", "DELETE":"Delete"
    }.get(e.split()[0], "Operation")) for e in endpoints], [4.9,1.6])

p("Appendix C. Glossary", "Heading 1")
table(["Term", "Meaning"], [
("Access code", "Administrator-generated code that grants or extends subscription entitlement."),
("Embedding", "Numeric representation of text used to retrieve semantically relevant source passages."),
("FAISS", "Vector-search library used for similarity retrieval."),
("LLM", "Large language model used for tutoring and content generation."),
("Mastery", "Recorded evidence of a learner’s performance for a subject/topic."),
("OTP", "Short-lived one-time password used to verify identity or a transaction."),
("PWA", "Progressive Web App; an installable web experience with offline/push capabilities."),
("RAG", "Retrieval-Augmented Generation: providing relevant source text to an LLM before it answers."),
("WASSCE", "West African Senior School Certificate Examination."),
("Webhook", "Provider-to-server callback used to report payment or other external events."),
])

add_footer()
doc.core_properties.title = "BroxStudies System Documentation"
doc.core_properties.subject = "Features, architecture, APIs, operations, and security"
doc.core_properties.author = "BroxStudies"
doc.core_properties.keywords = "BroxStudies, system documentation, features, API, architecture"
doc.save(OUT)
print(OUT.resolve())
